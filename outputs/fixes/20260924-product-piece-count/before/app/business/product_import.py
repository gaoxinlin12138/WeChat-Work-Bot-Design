"""Parse the two-row product archive used by the product administration UI."""

from __future__ import annotations

from decimal import Decimal, InvalidOperation
from io import BytesIO

from openpyxl import load_workbook

BASE_HEADERS = (
    "序号",
    "产品型号",
    "系列",
    "产品描述",
    "材质",
    "档位",
    "颜色",
    "特殊功能",
    "出货地",
    "套装价(元)",
    "主体价(元)",
    "成本价",
    "供应商",
    "主体供应商",
)
COMPONENT_HEADERS = (
    "说明书",
    "PVC管",
    "电话管",
    "顶喷",
    "手喷",
    "喷枪",
    "直管",
    "七字管",
    "中叉",
    "固定座",
    "装饰盖",
    "弯角",
    "六分橡皮",
    "膨胀螺丝",
    "生料带",
    "内盒",
    "外箱",
    "布袋",
    "珍珠棉带",
    "泡沫盒",
)
PRODUCT_HEADERS = (*BASE_HEADERS, *COMPONENT_HEADERS, "备注")
EMPTY_MARKERS = {"", "暂无", "无", "没有", "n/a", "N/A", "-", "/"}


def _text(value: object) -> str:
    if value is None:
        return ""
    text = str(value).strip()
    return "" if text in EMPTY_MARKERS else text


def _integer(value: object, *, row: int, field: str) -> int | None:
    text = _text(value)
    if not text:
        return None
    try:
        number = Decimal(text)
    except InvalidOperation as exc:
        raise ValueError(f"第 {row} 行“{field}”必须是整数") from exc
    if number != number.to_integral_value():
        raise ValueError(f"第 {row} 行“{field}”必须是整数")
    return int(number)


def _money(value: object, *, row: int, field: str) -> Decimal | None:
    text = _text(value).replace(",", "").replace("￥", "").replace("¥", "")
    if not text:
        return None
    try:
        number = Decimal(text).quantize(Decimal("0.01"))
    except InvalidOperation as exc:
        raise ValueError(f"第 {row} 行“{field}”不是有效金额") from exc
    if number < 0:
        raise ValueError(f"第 {row} 行“{field}”不能小于 0")
    return number


def _component_rows(
    product_values: tuple[object, ...],
    price_values: tuple[object, ...] | None,
    *,
    product_row: int,
    price_row: int | None,
) -> list[dict[str, object]]:
    rows: list[dict[str, object]] = []
    offset = len(BASE_HEADERS)
    for index, name in enumerate(COMPONENT_HEADERS):
        supplier = _text(product_values[offset + index])
        unit_price = (
            _money(price_values[offset + index], row=price_row or product_row, field=f"{name}价格")
            if price_values is not None
            else None
        )
        rows.append({"name": name, "supplier": supplier, "unit_price": unit_price})
    return rows


def parse_product_workbook(content: bytes) -> dict[str, object]:
    """Parse a product row followed by its optional component-price row."""
    try:
        workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
    except Exception as exc:
        raise ValueError("无法读取 Excel 文件，请确认文件未损坏") from exc
    if "产品档案" not in workbook.sheetnames:
        raise ValueError("Excel 缺少“产品档案”工作表")
    sheet = workbook["产品档案"]
    # Some valid XLSX writers omit the optional worksheet ``dimension`` tag.
    # In read-only mode openpyxl then reports ``max_row`` as ``None`` even
    # though cells can still be read normally.
    header_scan_limit = min(sheet.max_row or 20, 20)
    header_row = next(
        (
            row_number
            for row_number in range(1, header_scan_limit + 1)
            if _text(sheet.cell(row_number, 1).value) == "序号"
        ),
        None,
    )
    if header_row is None:
        raise ValueError("产品档案中未找到表头行")
    actual_headers = tuple(
        _text(sheet.cell(header_row, column).value) for column in range(1, len(PRODUCT_HEADERS) + 1)
    )
    if actual_headers != PRODUCT_HEADERS:
        raise ValueError("产品档案表头与系统模板不一致，请保留35列新版表头")

    rows = list(sheet.iter_rows(min_row=header_row + 1, values_only=True))
    records: list[dict[str, object]] = []
    warnings: list[dict[str, object]] = []
    index = 0
    while index < len(rows):
        source_row = header_row + 1 + index
        values = tuple(rows[index][: len(PRODUCT_HEADERS)])
        if not any(_text(value) for value in values):
            index += 1
            continue
        if _text(values[len(BASE_HEADERS) - 1]) == "价格":
            warnings.append(
                {"row": source_row, "field": "价格", "message": "未找到对应的产品行，未导入"}
            )
            index += 1
            continue
        model = _text(values[1])
        if not model:
            warnings.append(
                {"row": source_row, "field": "产品型号", "message": "缺少产品型号，未导入"}
            )
            index += 1
            continue
        price_values = None
        price_row = None
        if index + 1 < len(rows):
            candidate = tuple(rows[index + 1][: len(PRODUCT_HEADERS)])
            if _text(candidate[len(BASE_HEADERS) - 1]) == "价格":
                price_values, price_row = candidate, source_row + 1
        if price_values is None:
            warnings.append(
                {
                    "row": source_row,
                    "field": "配件价格",
                    "message": "缺少下一行配件价格，已按空价格导入",
                }
            )
        records.append(
            {
                "sequence_number": _integer(values[0], row=source_row, field="序号"),
                "model": model,
                "series": _text(values[2]),
                "product_description": _text(values[3]),
                "style": "",
                "material": _text(values[4]),
                "spray_modes": _text(values[5]),
                "color": _text(values[6]),
                "special_function": _text(values[7]),
                "shipping_origin": _text(values[8]),
                "set_price": _money(values[9], row=source_row, field="套装价(元)"),
                "body_price": _money(values[10], row=source_row, field="主体价(元)"),
                "cost_price": _money(values[11], row=source_row, field="成本价"),
                "supplier": _text(values[12]),
                "body_supplier": _text(values[13]),
                "components": _component_rows(
                    values, price_values, product_row=source_row, price_row=price_row
                ),
                "notes": _text(values[len(PRODUCT_HEADERS) - 1]),
                "source_row": source_row,
                "status": "draft",
            }
        )
        index += 2 if price_values is not None else 1
    if not records:
        raise ValueError("产品档案中没有可导入的有效记录")
    return {"records": records, "warnings": warnings}

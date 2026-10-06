"""Parse the supported product archive layouts used by product administration."""

from __future__ import annotations

from decimal import Decimal, InvalidOperation
from io import BytesIO

from openpyxl import load_workbook

BASE_HEADERS = (
    "序号",
    "产品型号",
    "系列",
    "产品描述",
    "材质",
    "档位",
    "颜色",
    "特殊功能",
    "出货地",
    "套装价",
    "成本价",
    "主体价",
    "供应商",
    "主体供应商",
)
COMPONENT_HEADERS = (
    "说明书",
    "PVC管",
    "电话管",
    "顶喷",
    "手喷",
    "喷枪",
    "淋浴杆",
    "中叉",
    "固定座",
    "装饰盖*2",
    "弯角*2",
    "橡皮*4",
    "止水阀*2",
    "生料带",
    "内盒",
    "外箱",
    "布袋",
    "珍珠棉袋",
    "泡沫盒",
)
PRODUCT_HEADERS = (*BASE_HEADERS, *COMPONENT_HEADERS, "备注")
EMPTY_MARKERS = {"", "暂无", "无", "没有", "n/a", "N/A", "-", "/"}
_HEADER_ALIASES = {
    "套装价(元)": "套装价",
    "套装价（元）": "套装价",
    "主体价(元)": "主体价",
    "主体价（元）": "主体价",
    "成本价(元)": "成本价",
    "成本价（元）": "成本价",
}


def _text(value: object) -> str:
    if value is None:
        return ""
    text = str(value).strip()
    return "" if text in EMPTY_MARKERS else text


def _header(value: object) -> str:
    text = _text(value)
    return _HEADER_ALIASES.get(text, text)


def _integer(value: object, *, row: int, field: str) -> int | None:
    text = _text(value)
    if not text:
        return None
    try:
        number = Decimal(text)
    except InvalidOperation as exc:
        raise ValueError(f"第 {row} 行“{field}”必须是整数") from exc
    if not number.is_finite() or number != number.to_integral_value():
        raise ValueError(f"第 {row} 行“{field}”必须是整数")
    return int(number)


def _piece_count(value: object, *, row: int) -> int | None:
    count = _integer(value, row=row, field="件数")
    if count is not None and not 0 <= count <= 2147483647:
        raise ValueError(f"第 {row} 行“件数”必须在 0 到 2147483647 之间")
    return count


def _money(value: object, *, row: int, field: str) -> Decimal | None:
    text = _text(value).replace(",", "").replace("￥", "").replace("¥", "")
    if not text:
        return None
    try:
        number = Decimal(text).quantize(Decimal("0.01"))
    except InvalidOperation as exc:
        raise ValueError(f"第 {row} 行“{field}”不是有效金额") from exc
    if not number.is_finite():
        raise ValueError(f"第 {row} 行“{field}”不是有效金额")
    if number < 0:
        raise ValueError(f"第 {row} 行“{field}”不能小于 0")
    return number


def _money_if_numeric(value: object, *, row: int, field: str) -> Decimal | None:
    if value is None or isinstance(value, bool):
        return None
    if isinstance(value, (int, float, Decimal)):
        return _money(value, row=row, field=field)
    text = _text(value).replace(",", "").replace("￥", "").replace("¥", "")
    if not text:
        return None
    try:
        Decimal(text)
    except InvalidOperation:
        return None
    return _money(value, row=row, field=field)


def _compact_components(
    product_values: tuple[object, ...],
    price_values: tuple[object, ...] | None,
    *,
    product_row: int,
    price_row: int | None,
) -> list[dict[str, object]]:
    rows: list[dict[str, object]] = []
    offset = len(BASE_HEADERS)
    for index, name in enumerate(COMPONENT_HEADERS):
        rows.append(
            {
                "name": name,
                "supplier": _text(product_values[offset + index]),
                "unit_price": (
                    _money(
                        price_values[offset + index],
                        row=price_row or product_row,
                        field=f"{name}价格",
                    )
                    if price_values is not None
                    else None
                ),
            }
        )
    return rows


def _paired_components(
    values: tuple[object, ...], *, product_row: int
) -> list[dict[str, object]]:
    rows: list[dict[str, object]] = []
    offset = len(BASE_HEADERS)
    for index, name in enumerate(COMPONENT_HEADERS):
        pair = values[offset + index * 2 : offset + index * 2 + 2]
        supplier_parts: list[str] = []
        unit_price: Decimal | None = None
        for value in pair:
            numeric = _money_if_numeric(value, row=product_row, field=f"{name}价格")
            if numeric is not None:
                if unit_price is not None:
                    raise ValueError(f"第 {product_row} 行“{name}”只能填写一个采购价")
                unit_price = numeric
            elif text := _text(value):
                supplier_parts.append(text)
        rows.append(
            {
                "name": name,
                "supplier": " / ".join(dict.fromkeys(supplier_parts)),
                "unit_price": unit_price,
            }
        )
    return rows


def _without_column(row: tuple[object, ...], column: int | None) -> tuple[object, ...]:
    if column is None:
        return row
    return tuple(value for index, value in enumerate(row) if index != column)


def _detect_layout(headers: tuple[str, ...]) -> str:
    compact = tuple(headers[: len(PRODUCT_HEADERS)])
    if compact == PRODUCT_HEADERS and not any(headers[len(PRODUCT_HEADERS) :]):
        return "compact"
    paired_headers = [*BASE_HEADERS]
    for name in COMPONENT_HEADERS:
        paired_headers.extend((name, ""))
    paired_headers.append("备注")
    if (
        tuple(headers[: len(paired_headers)]) == tuple(paired_headers)
        and not any(headers[len(paired_headers) :])
    ):
        return "paired"
    raise ValueError(
        "产品档案表头与系统字段不一致，请使用当前 34 项产品字段；"
        "“件数”列可选且可放在任意位置"
    )


def parse_product_workbook(content: bytes) -> dict[str, object]:
    """Parse compact rows or paired supplier/price component columns."""
    try:
        workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
    except Exception as exc:
        raise ValueError("无法读取 Excel 文件，请确认文件未损坏") from exc
    if "产品档案" not in workbook.sheetnames:
        raise ValueError("Excel 缺少“产品档案”工作表")
    sheet = workbook["产品档案"]
    header_scan_limit = min(sheet.max_row or 20, 20)
    header_row = next(
        (
            row_number
            for row_number in range(1, header_scan_limit + 1)
            if _header(sheet.cell(row_number, 1).value) == "序号"
            or (
                _header(sheet.cell(row_number, 1).value) == "件数"
                and _header(sheet.cell(row_number, 2).value) == "序号"
            )
        ),
        None,
    )
    if header_row is None:
        raise ValueError("产品档案中未找到表头行")
    raw_headers = tuple(
        next(sheet.iter_rows(min_row=header_row, max_row=header_row, values_only=True))
    )
    normalized_headers = tuple(_header(value) for value in raw_headers)
    if normalized_headers.count("件数") > 1:
        raise ValueError("产品档案中只能有一列“件数”")
    piece_count_column = (
        normalized_headers.index("件数") if "件数" in normalized_headers else None
    )
    headers = _without_column(normalized_headers, piece_count_column)
    layout = _detect_layout(headers)
    rows = list(sheet.iter_rows(min_row=header_row + 1, values_only=True))
    records: list[dict[str, object]] = []
    model_rows: dict[str, int] = {}
    warnings: list[dict[str, object]] = []
    index = 0
    while index < len(rows):
        source_row = header_row + 1 + index
        raw_values = rows[index]
        values = _without_column(raw_values, piece_count_column)
        if not any(_text(value) for value in values):
            index += 1
            continue
        if layout == "compact" and _text(values[len(BASE_HEADERS) - 1]) == "价格":
            warnings.append(
                {"row": source_row, "field": "价格", "message": "未找到对应的产品行，未导入"}
            )
            index += 1
            continue
        model = _text(values[1])
        if not model:
            warnings.append(
                {"row": source_row, "field": "产品型号", "message": "缺少产品型号，未导入"}
            )
            index += 1
            continue
        model_key = model.casefold()
        if model_key in model_rows:
            raise ValueError(
                f"产品型号“{model}”在第 {model_rows[model_key]} 行和第 {source_row} 行重复；"
                "产品型号必须唯一"
            )
        model_rows[model_key] = source_row
        price_values = None
        price_row = None
        if layout == "compact" and index + 1 < len(rows):
            candidate = _without_column(rows[index + 1], piece_count_column)
            if _text(candidate[len(BASE_HEADERS) - 1]) == "价格":
                price_values, price_row = candidate, source_row + 1
        if layout == "compact" and price_values is None:
            warnings.append(
                {
                    "row": source_row,
                    "field": "配件价格",
                    "message": "缺少下一行配件价格，已按空价格导入",
                }
            )
        components = (
            _compact_components(
                values,
                price_values,
                product_row=source_row,
                price_row=price_row,
            )
            if layout == "compact"
            else _paired_components(values, product_row=source_row)
        )
        notes_column = (
            len(PRODUCT_HEADERS) - 1
            if layout == "compact"
            else len(BASE_HEADERS) + len(COMPONENT_HEADERS) * 2
        )
        records.append(
            {
                "sequence_number": _integer(values[0], row=source_row, field="序号"),
                "model": model,
                "series": _text(values[2]),
                "product_description": _text(values[3]),
                "material": _text(values[4]),
                "spray_modes": _text(values[5]),
                "color": _text(values[6]),
                "special_function": _text(values[7]),
                "shipping_origin": _text(values[8]),
                "set_price": _money(values[9], row=source_row, field="套装价"),
                "cost_price": _money(values[10], row=source_row, field="成本价"),
                "body_price": _money(values[11], row=source_row, field="主体价"),
                "piece_count": (
                    _piece_count(raw_values[piece_count_column], row=source_row)
                    if piece_count_column is not None and piece_count_column < len(raw_values)
                    else None
                ),
                "supplier": _text(values[12]),
                "body_supplier": _text(values[13]),
                "components": components,
                "notes": _text(values[notes_column]),
                "source_row": source_row,
                "status": "draft",
            }
        )
        index += 2 if price_values is not None else 1
    workbook.close()
    if not records:
        raise ValueError("产品档案中没有可导入的有效记录")
    return {"records": records, "warnings": warnings}

from io import BytesIO

import pytest
from fastapi.testclient import TestClient
from openpyxl import Workbook

import app.business.router as business_router
from app.business.intent import Intent, IntentDecision
from app.business.media_files import delete_uploaded_media_files
from app.business.planner import ProductFilter, ProductQuery
from app.business.product_import import BASE_HEADERS, PRODUCT_HEADERS, parse_product_workbook
from app.business.product_store import ProductModelConflictError
from app.business.service import BusinessAgent
from app.business.store import BusinessStore
from app.main import app
from scripts.cleanup_orphan_media import find_orphan_media_files


def workbook_bytes(rows: list[list[object]]) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "产品档案"
    sheet.append(["花洒产品档案"])
    sheet.append(list(PRODUCT_HEADERS))
    for row in rows:
        sheet.append(row)
    workbook.create_sheet("说明")["A1"] = "这里的内容只是资料说明"
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def sample_row(model: str = "LTA-822") -> list[object]:
    return [
        1,
        model,
        "太空铝系列",
        "套装花洒",
        "全铝材质",
        4,
        "灰色",
        "恒温",
        "浙江宁波",
        170,
        140,
        110,
        "供应商甲",
        "主体供应商乙",
        "说明书供应商",
        "PVC供应商",
        "电话管供应商",
        "顶喷供应商",
        "手喷供应商",
        "喷枪供应商",
        "淋浴杆供应商",
        "中叉供应商",
        "固定座供应商",
        "装饰盖*2供应商",
        "弯角*2供应商",
        "橡皮供应商",
        "止水阀*2供应商",
        "生料带供应商",
        "内盒供应商",
        "外箱供应商",
        "布袋供应商",
        "珍珠棉供应商",
        "泡沫盒供应商",
        "备注",
    ]


def sample_price_row() -> list[object]:
    row = [None] * len(PRODUCT_HEADERS)
    row[len(BASE_HEADERS) - 1] = "价格"
    row[len(BASE_HEADERS)] = 0.26
    row[len(BASE_HEADERS) + 1] = 3.6
    row[len(BASE_HEADERS) + 2] = 2.2
    return row


def paired_workbook_bytes() -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "产品档案"
    headers = list(BASE_HEADERS)
    for name in PRODUCT_HEADERS[len(BASE_HEADERS):-1]:
        headers.extend((name, None))
    headers.append("备注")
    row = sample_row()[:len(BASE_HEADERS)]
    row.extend(["说明书供应商", 0.26])
    for _ in PRODUCT_HEADERS[len(BASE_HEADERS) + 1:-1]:
        row.extend((None, None))
    row.append("双列备注")
    sheet.append(headers)
    sheet.append(row)
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def test_product_workbook_maps_product_and_component_columns() -> None:
    parsed = parse_product_workbook(workbook_bytes([sample_row(), sample_price_row()]))
    assert len(parsed["records"]) == 1
    record = parsed["records"][0]
    assert record["model"] == "LTA-822"
    assert record["set_price"] == 170
    assert record["cost_price"] == 140
    assert record["body_price"] == 110
    assert "main_image_file" not in record
    assert "gallery_image_files" not in record
    assert "configuration_image_file" not in record
    assert "video_file" not in record
    assert record["body_supplier"] == "主体供应商乙"
    assert float(record["components"][1]["unit_price"]) == 3.6
    assert record["source_row"] == 3


def test_product_workbook_maps_paired_supplier_and_price_columns() -> None:
    record = parse_product_workbook(paired_workbook_bytes())["records"][0]
    assert record["set_price"] == 170
    assert record["cost_price"] == 140
    assert record["body_price"] == 110
    assert record["components"][0]["name"] == "说明书"
    assert record["components"][0]["supplier"] == "说明书供应商"
    assert float(record["components"][0]["unit_price"]) == 0.26
    assert record["notes"] == "双列备注"


def test_product_workbook_rejects_duplicate_models() -> None:
    with pytest.raises(ValueError, match="产品型号必须唯一"):
        parse_product_workbook(
            workbook_bytes([sample_row(), sample_price_row(), sample_row(), sample_price_row()])
        )


class ProductApiStore:
    def __init__(self) -> None:
        self.products: list[dict[str, object]] = []

    def import_product_records(self, records):
        self.products = [
            {"id": str(index), "media": [], **item} for index, item in enumerate(records)
        ]
        return self.products

    def list_product_records(self, include_deleted=False):
        return self.products


class ProductDeleteResult:
    def __init__(self, rows):
        self.rows = rows

    def fetchall(self):
        return self.rows


class ProductDeleteConnection:
    def __init__(self, rows):
        self.rows = rows
        self.statements = []

    def __enter__(self):
        return self

    def __exit__(self, exc_type, exc, traceback):
        return False

    def execute(self, statement, params=()):
        self.statements.append((statement, params))
        return ProductDeleteResult(self.rows)


def test_product_import_endpoint_uses_new_archive_route(monkeypatch) -> None:
    store = ProductApiStore()
    monkeypatch.setattr(business_router, "get_business_store", lambda: store)
    response = TestClient(app).post(
        "/business/products/import-xlsx",
        files={
            "file": (
                "花洒产品档案.xlsx",
                workbook_bytes([sample_row()]),
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        },
    )
    assert response.status_code == 200
    assert response.json()["count"] == 1
    assert response.json()["products"][0]["model"] == "LTA-822"


def test_bulk_delete_product_store_deduplicates_and_returns_changed_count(monkeypatch) -> None:
    connection = ProductDeleteConnection([{"id": "product-1"}, {"id": "product-2"}])
    store = BusinessStore("postgresql://unused")
    monkeypatch.setattr(store, "_connect", lambda: connection)

    assert store.delete_product_records(["product-1", "product-1", "product-2"]) == 2
    statement, params = connection.statements[0]
    assert "status='deleted'" in statement
    assert "status<>'deleted'" in statement
    assert "RETURNING id::text" in statement
    assert params == (["product-1", "product-2"],)


def test_product_video_above_safe_limit_is_rejected_before_storage(monkeypatch) -> None:
    monkeypatch.setattr(business_router, "WECOM_VIDEO_SAFE_MAX_BYTES", 5)

    response = TestClient(app).post(
        "/business/products/product-1/media",
        files={"file": ("demo.mp4", b"123456", "video/mp4")},
        data={"role": "product_video", "sort_order": "0"},
    )

    assert response.status_code == 413
    assert "9MB" in response.json()["detail"]


def test_product_image_above_send_limit_is_rejected_before_storage(monkeypatch) -> None:
    monkeypatch.setattr(business_router, "WECOM_IMAGE_SAFE_MAX_BYTES", 5)

    response = TestClient(app).post(
        "/business/products/product-1/media",
        files={"file": ("demo.png", b"123456", "image/png")},
        data={"role": "product_image", "sort_order": "0"},
    )

    assert response.status_code == 413
    assert "10MB" in response.json()["detail"]


def test_media_cleanup_handles_current_schema_without_source_type(tmp_path) -> None:
    upload_dir = tmp_path / "media"
    upload_dir.mkdir()
    file_path = upload_dir / "old.mp4"
    file_path.write_bytes(b"old")

    assert delete_uploaded_media_files(
        [{"storage_path": str(file_path)}], str(upload_dir)
    ) == 1
    assert not file_path.exists()


def test_orphan_media_scan_keeps_referenced_files(tmp_path) -> None:
    kept = tmp_path / "kept.png"
    orphan = tmp_path / "orphan.mp4"
    kept.write_bytes(b"image")
    orphan.write_bytes(b"video")

    assert find_orphan_media_files(str(tmp_path), [str(kept)]) == [orphan]


def test_database_rejects_duplicate_product_models(isolated_database_url) -> None:
    store = BusinessStore(isolated_database_url)
    store.initialize()
    values = parse_product_workbook(workbook_bytes([sample_row()]))["records"][0]
    first = store.create_product_record(values)
    with pytest.raises(ProductModelConflictError, match="全局唯一"):
        store.create_product_record({**values, "model": " lta-822 "})
    assert [item["model"] for item in store.list_product_records()] == [first["model"]]


def test_bulk_delete_products_deduplicates_and_only_trashes_live_records(
    isolated_database_url,
) -> None:
    store = BusinessStore(isolated_database_url)
    store.initialize()
    values = parse_product_workbook(workbook_bytes([sample_row()]))["records"][0]
    first = store.create_product_record(values)
    second = store.create_product_record({**values, "model": "LTA-823"})

    assert store.delete_product_records([first["id"], first["id"], second["id"]]) == 2
    assert store.delete_product_records([first["id"], second["id"]]) == 0
    statuses = {
        item["id"]: item["status"] for item in store.list_product_records(include_deleted=True)
    }
    assert statuses[first["id"]] == "deleted"
    assert statuses[second["id"]] == "deleted"


def test_product_video_role_keeps_multiple_files(
    isolated_database_url,
) -> None:
    store = BusinessStore(isolated_database_url)
    store.initialize()
    values = parse_product_workbook(workbook_bytes([sample_row()]))["records"][0]
    product = store.create_product_record(values)

    first = store.create_product_media(
        product["id"],
        role="product_video",
        original_name="old.mp4",
        stored_path="data/product_media/old.mp4",
        media_type="video",
        content_hash="a" * 64,
    )
    second = store.create_product_media(
        product["id"],
        role="product_video",
        original_name="new.mp4",
        stored_path="data/product_media/new.mp4",
        media_type="video",
        content_hash="b" * 64,
    )

    assert first["replaced_media"] == []
    assert second["replaced_media"] == []
    assert [
        item["original_name"]
        for item in store.get_product_record(product["id"])["media"]
        if item["role"] == "product_video"
    ] == ["old.mp4", "new.mp4"]


def test_product_cost_fields_require_explicit_permission(isolated_database_url) -> None:
    store = BusinessStore(isolated_database_url)
    store.initialize()
    values = parse_product_workbook(
        workbook_bytes([sample_row(), sample_price_row()])
    )["records"][0]
    created = store.create_product_record(values)
    store.set_product_status(created["id"], "active")

    search_text = store.build_product_search_text(created["id"])
    product = store.get_active_product(created["id"])
    agent = BusinessAgent(store)
    public_reply = agent._format_product_detail(product, [], "配件有吗，都是些什么")
    denied_reply = agent._format_product_detail(product, [], "成本价是多少")
    sales_reply = agent._format_product_detail(
        product,
        [],
        "配件采购价是多少",
        can_view_cost=True,
    )
    sales_overview = agent._format_product_detail(product, [], can_view_cost=True)

    assert "配件明细：说明书｜供应商：说明书供应商" in search_text
    assert "成本价" not in search_text
    assert "采购价" not in search_text
    assert "配件明细（共19项）" in public_reply.content
    assert "采购价" not in public_reply.content
    assert denied_reply.content == "你当前没有查看成本价和采购价的权限。"
    assert "1. 说明书｜供应商：说明书供应商｜采购价：0.26元" in sales_reply.content
    assert "2. PVC管｜供应商：PVC供应商｜采购价：3.6元" in sales_reply.content
    assert "成本价：140.00元" in sales_overview.content


def test_product_catalog_uses_compact_model_rows(isolated_database_url) -> None:
    store = BusinessStore(isolated_database_url)
    store.initialize()
    created = store.create_product_record(
        {
            "model": "LTA-822",
            "series": "太空铝花洒",
            "product_description": "按键四档带电机、氛围灯套装花洒",
            "material": "全铝主体",
            "spray_modes": "4",
            "color": "灰玻璃·镀灰",
            "set_price": 170,
            "body_price": 110,
            "piece_count": 150,
        }
    )
    store.set_product_status(created["id"], "active")

    reply = BusinessAgent(store)._format_product_catalog(
        [store.get_active_product(created["id"])],
    )

    assert "已发布产品共展示 1 款" in reply.content
    assert "1. LTA-822" in reply.content
    assert "库存：150件" in reply.content
    assert "回复“查看+型号”（例如“查看LTA-822”）可获取完整参数、配件和媒体资料。" in reply.content
    assert "配件明细" not in reply.content


def test_product_shorthand_price_query_lists_every_matching_model(
    isolated_database_url,
) -> None:
    store = BusinessStore(isolated_database_url)
    store.initialize()
    for model in ("LTA-822", "MTA-822", "NTA-822", "OTA-822"):
        created = store.create_product_record({"model": model, "set_price": 170})
        store.set_product_status(created["id"], "active")

    query = "822配件价格"
    reply = BusinessAgent(store)._answer_product(
        "employee-1",
        query,
        query,
        IntentDecision(Intent.PRODUCT, 0.99, "查询产品配件价格", query),
    )

    assert "822系列目前有 4 款" in reply.content
    assert "1. LTA-822" in reply.content
    assert "4. OTA-822" in reply.content
    assert reply.content.count("套装价：170.00元") == 4
    assert "请补充具体型号" not in reply.content


def test_product_shorthand_is_not_limited_to_four_digits(isolated_database_url) -> None:
    store = BusinessStore(isolated_database_url)
    store.initialize()
    for model in ("FAMILY-A72-GT", "FAMILY-A72-NB"):
        created = store.create_product_record({"model": model, "set_price": 188})
        store.set_product_status(created["id"], "active")

    query = "A72价格多少"
    reply = BusinessAgent(store)._answer_product(
        "employee-1",
        query,
        query,
        IntentDecision(Intent.PRODUCT, 0.99, "查询型号简称价格", query),
    )

    assert "a72系列目前有 2 款" in reply.content.lower()
    assert "FAMILY-A72-GT" in reply.content
    assert "FAMILY-A72-NB" in reply.content


def test_product_family_full_information_is_paginated(isolated_database_url) -> None:
    store = BusinessStore(isolated_database_url)
    store.initialize()
    for model, color in (("LY0653GDT", "灰色"), ("LY0653NBT", "奶白")):
        created = store.create_product_record(
            {
                "model": model,
                "color": color,
                "set_price": 200,
                "components": [{"name": "说明书", "supplier": "供应商甲"}],
            }
        )
        store.set_product_status(created["id"], "active")
    agent = BusinessAgent(store)
    query = "653全部信息"
    decision = IntentDecision(Intent.PRODUCT, 0.99, "查询系列完整资料", query)

    first = agent._answer_product("employee-1", query, query, decision)
    second = agent._answer_product("employee-1", "下一页", "下一页", decision)

    assert "653系列共 2 款，完整资料第 1/2 页" in first.content
    assert "产品型号：LY0653GDT" in first.content
    assert "产品型号：LY0653NBT" not in first.content
    assert "配件明细（共1项）" in first.content
    assert "653系列共 2 款，完整资料第 2/2 页" in second.content
    assert "产品型号：LY0653NBT" in second.content
    assert "产品型号：LY0653GDT" not in second.content


def test_product_query_without_model_asks_for_model(isolated_database_url) -> None:
    store = BusinessStore(isolated_database_url)
    store.initialize()
    created = store.create_product_record({"model": "LTA-822", "set_price": 170})
    store.set_product_status(created["id"], "active")

    query = "配件价格"
    reply = BusinessAgent(store)._answer_product(
        "employee-1",
        query,
        query,
        IntentDecision(Intent.PRODUCT, 0.99, "缺少产品型号", query),
    )

    assert reply.content.startswith("请先提供具体产品型号")
    assert "LTA-822" not in reply.content


def test_product_catalog_is_paginated_and_supports_next_page(isolated_database_url) -> None:
    store = BusinessStore(isolated_database_url)
    store.initialize()
    for index in range(7):
        created = store.create_product_record({"model": f"CAT-{index + 1:03d}"})
        store.set_product_status(created["id"], "active")
    agent = BusinessAgent(store)
    decision = IntentDecision(Intent.PRODUCT, 0.99, "查看产品目录", "查看产品目录")

    first_page = agent._answer_product(
        "employee-1",
        "查看产品目录",
        "查看产品目录",
        decision,
        product_query=ProductQuery(operation="list"),
    )
    assert "第 1/2 页" in first_page.content
    assert "1. CAT-001" in first_page.content
    assert "5. CAT-005" in first_page.content
    assert "6. CAT-006" not in first_page.content

    next_page = agent._answer_product("employee-1", "下一页", "下一页", decision)
    assert "第 2/2 页" in next_page.content
    assert "6. CAT-006" in next_page.content
    assert "7. CAT-007" in next_page.content
    assert "1. CAT-001" not in next_page.content


def test_special_function_catalog_starts_new_filtered_pagination(
    isolated_database_url,
) -> None:
    store = BusinessStore(isolated_database_url)
    store.initialize()
    for index in range(6):
        created = store.create_product_record(
            {
                "model": f"SPECIAL-{index + 1:03d}",
                "special_function": "恒温、防烫" if index % 2 == 0 else "增压",
            }
        )
        store.set_product_status(created["id"], "active")
    for index, value in enumerate(("", "未登记"), 1):
        created = store.create_product_record(
            {"model": f"EMPTY-{index:03d}", "special_function": value}
        )
        store.set_product_status(created["id"], "active")

    agent = BusinessAgent(store)
    decision = IntentDecision(Intent.PRODUCT, 0.99, "查看产品目录", "查看产品目录")
    agent._answer_product(
        "employee-1",
        "查看产品目录",
        "查看产品目录",
        decision,
        product_query=ProductQuery(operation="list"),
    )
    previous_second_page = agent._answer_product(
        "employee-1", "下一页", "下一页", decision
    )
    assert "第 2/2 页" in previous_second_page.content

    query = ProductQuery(
        operation="list",
        fields=("model", "special_function"),
        filters=(ProductFilter("special_function", "is_present", ""),),
    )
    first_page = agent._answer_product(
        "employee-1",
        "哪些产品含有特殊功能",
        "哪些产品含有特殊功能",
        decision,
        product_query=query,
    )

    assert "找到 6 款已登记特殊功能的产品，第 1/2 页" in first_page.content
    assert "特殊功能：" in first_page.content
    assert "EMPTY-" not in first_page.content

    next_page = agent._answer_product("employee-1", "下一页", "下一页", decision)
    assert "找到 6 款已登记特殊功能的产品，第 2/2 页" in next_page.content
    assert "特殊功能：" in next_page.content
    assert "EMPTY-" not in next_page.content


def test_special_function_catalog_filters_specific_value_and_reports_empty(
    isolated_database_url,
) -> None:
    store = BusinessStore(isolated_database_url)
    store.initialize()
    for model, special_function in (
        ("THERMO-001", "恒温、防烫"),
        ("BOOST-001", "增压"),
        ("PLAIN-001", ""),
    ):
        created = store.create_product_record(
            {"model": model, "special_function": special_function}
        )
        store.set_product_status(created["id"], "active")
    agent = BusinessAgent(store)
    decision = IntentDecision(Intent.PRODUCT, 0.99, "筛选特殊功能", "哪些产品支持恒温")

    matched = agent._answer_product(
        "employee-1",
        "哪些产品支持恒温",
        "哪些产品支持恒温",
        decision,
        product_query=ProductQuery(
            operation="list",
            fields=("model", "special_function"),
            filters=(ProductFilter("special_function", "contains", "恒温"),),
        ),
    )
    missing = agent._answer_product(
        "employee-1",
        "哪些产品支持语音控制",
        "哪些产品支持语音控制",
        decision,
        product_query=ProductQuery(
            operation="list",
            fields=("model", "special_function"),
            filters=(ProductFilter("special_function", "contains", "语音控制"),),
        ),
    )

    assert "找到 1 款特殊功能包含“恒温”的产品" in matched.content
    assert "1. THERMO-001" in matched.content
    assert "特殊功能：恒温、防烫" in matched.content
    assert "BOOST-001" not in matched.content
    assert missing.content == "当前没有特殊功能包含“语音控制”的已发布产品。"


def test_component_total_uses_quantity_suffix_and_requires_price_permission(
    isolated_database_url,
) -> None:
    store = BusinessStore(isolated_database_url)
    store.initialize()
    created = store.create_product_record(
        {
            "model": "TOTAL-001",
            "components": [
                {"name": "装饰盖*2", "supplier": "供应商甲", "unit_price": 3},
                {"name": "弯角", "supplier": "供应商乙", "unit_price": 4.5},
            ],
        }
    )
    store.set_product_status(created["id"], "active")
    product = store.get_active_product(created["id"])
    agent = BusinessAgent(store)

    permitted = agent._format_product_detail(
        product,
        [],
        "配件总价是多少",
        can_view_cost=True,
        requested_fields=("components",),
    )
    denied = agent._format_product_detail(
        product,
        [],
        "配件总价是多少",
        can_view_cost=False,
        requested_fields=("components",),
    )

    assert "配件采购总价：10.50元" in permitted.content
    assert "装饰盖*2｜供应商：供应商甲｜采购价：3元" not in permitted.content
    assert "没有查看成本价和采购价的权限" in denied.content
    assert "配件采购总价：10.50元" not in denied.content


def test_product_catalog_exclusion_returns_direct_empty_result(isolated_database_url) -> None:
    store = BusinessStore(isolated_database_url)
    store.initialize()
    created = store.create_product_record({"model": "LTA-822", "set_price": 170})
    store.set_product_status(created["id"], "active")
    query = "除了822，我们现在还有其他产品吗？"

    reply = BusinessAgent(store)._answer_product(
        "employee-1",
        query,
        query,
        IntentDecision(Intent.PRODUCT, 0.99, "产品排除查询", query),
        product_query=ProductQuery(operation="list", exclude_models=("822",)),
    )

    assert reply.content == "目前已发布产品只有 LTA-822，除此之外暂无其他产品。"
    assert len(reply.sources) == 1


def test_product_catalog_exclusion_keeps_only_other_products(isolated_database_url) -> None:
    store = BusinessStore(isolated_database_url)
    store.initialize()
    for model in ("LTA-822", "LTA-830"):
        created = store.create_product_record({"model": model, "set_price": 170})
        store.set_product_status(created["id"], "active")
    query = "除了LTA-822，我们现在还有其他产品吗？"

    reply = BusinessAgent(store)._answer_product(
        "employee-1",
        query,
        query,
        IntentDecision(Intent.PRODUCT, 0.99, "产品排除查询", query),
        product_query=ProductQuery(operation="list", exclude_models=("LTA-822",)),
    )

    assert reply.content.startswith("除 LTA-822 外，已发布产品还有 1 款：")
    assert "1. LTA-830" in reply.content
    assert "1. LTA-822" not in reply.content
    assert "查看LTA-830" in reply.content


def test_product_search_accepts_unique_model_suffix(isolated_database_url) -> None:
    store = BusinessStore(isolated_database_url)
    store.initialize()
    created = store.create_product_record({"model": "LTA-822", "set_price": 170})
    store.set_product_status(created["id"], "active")

    results = store.search_products("822图片发我")

    assert len(results) == 1
    assert "产品型号：LTA-822" in results[0].content


def test_product_exclusion_rejects_ambiguous_model_suffix(isolated_database_url) -> None:
    store = BusinessStore(isolated_database_url)
    store.initialize()
    for model in ("LTA-822", "MTA-822"):
        created = store.create_product_record({"model": model})
        store.set_product_status(created["id"], "active")
    query = "除了822，我们还有其他产品吗？"

    reply = BusinessAgent(store)._answer_product(
        "employee-1",
        query,
        query,
        IntentDecision(Intent.PRODUCT, 0.99, "产品排除查询", query),
        product_query=ProductQuery(operation="list", exclude_models=("822",)),
    )

    assert "可对应多个产品" in reply.content
    assert "LTA-822" in reply.content
    assert "MTA-822" in reply.content
    assert "请补充完整型号" in reply.content


def test_product_combined_fields_keep_public_and_sensitive_boundaries(
    isolated_database_url,
) -> None:
    store = BusinessStore(isolated_database_url)
    store.initialize()
    created = store.create_product_record(
        {
            "model": "COMBINED-001",
            "material": "全铝",
            "spray_modes": "三档",
            "set_price": 100,
            "body_price": 80,
            "cost_price": 60,
        }
    )
    store.set_product_status(created["id"], "active")
    product = store.get_active_product(created["id"])
    agent = BusinessAgent(store)

    public = agent._format_product_detail(product, [], "主体多少钱和材质")
    assert "主体价：80.00元" in public.content
    assert "材质：全铝" in public.content
    assert "套装价" not in public.content

    denied = agent._format_product_detail(product, [], "主体多少钱和成本价")
    assert "主体价：80.00元" in denied.content
    assert "没有查看成本价" in denied.content
    assert "成本价：60.00元" not in denied.content

    empty = store.create_product_record({"model": "EMPTY-PRICE-001"})
    store.set_product_status(empty["id"], "active")
    assert "资料中未登记" in (store.search_products("EMPTY-PRICE-001")[0].content)
    assert "None元" not in store.build_product_search_text(empty["id"])


def test_structured_fields_do_not_depend_on_query_keywords(isolated_database_url) -> None:
    store = BusinessStore(isolated_database_url)
    store.initialize()
    created = store.create_product_record({"model": "FIELD-001", "set_price": 188})
    store.set_product_status(created["id"], "active")
    product = store.get_active_product(created["id"])

    reply = BusinessAgent(store)._format_product_detail(
        product,
        [],
        "这个预算合适吗",
        requested_fields=("set_price",),
    )

    assert reply.content == "套装价：188.00元"


def test_deleted_product_still_reserves_unique_model(isolated_database_url) -> None:
    store = BusinessStore(isolated_database_url)
    store.initialize()
    created = store.create_product_record({"model": "RESERVED-001"})
    assert store.delete_product_record(created["id"])

    with pytest.raises(ProductModelConflictError, match="全局唯一"):
        store.create_product_record({"model": "reserved-001"})


def test_update_cannot_take_another_product_model(isolated_database_url) -> None:
    store = BusinessStore(isolated_database_url)
    store.initialize()
    first = store.create_product_record({"model": "UPDATE-001"})
    second = store.create_product_record({"model": "UPDATE-002"})

    with pytest.raises(ProductModelConflictError, match="全局唯一"):
        store.update_product_record(second["id"], {"model": " update-001 "})

    assert store.get_product_record(first["id"])["model"] == "UPDATE-001"
    assert store.get_product_record(second["id"])["model"] == "UPDATE-002"
